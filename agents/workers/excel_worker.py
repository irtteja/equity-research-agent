"""Excel / CSV extraction worker."""
import csv, io


def extract(file_path: str) -> dict:
    ext = file_path.rsplit(".", 1)[-1].lower()
    if ext in ("xlsx", "xls", "xlsm"):
        return _from_excel(file_path)
    return _from_csv(file_path)


def _from_excel(file_path: str) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheets = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                rows.append([str(c) if c is not None else "" for c in row])
        sheets[name] = rows
    wb.close()
    # flatten to text for uniform storage
    lines = []
    for sheet, rows in sheets.items():
        lines.append(f"=== Sheet: {sheet} ===")
        for r in rows:
            lines.append("\t".join(r))
    text = "\n".join(lines)
    return {"text": text, "pages": len(wb.sheetnames), "method": "openpyxl", "sheets": sheets}


def _from_csv(file_path: str) -> dict:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()
    return {"text": content, "pages": 1, "method": "csv"}
